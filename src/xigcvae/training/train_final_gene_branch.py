#!/usr/bin/env python3
"""
Train final gene-expression branch using the selected Top 2000 limma-ranked probes.

Outputs:
  outputs/gene_branch/final_gene_branch_model.pt
  outputs/gene_branch/final_gene_branch_scaler.pkl
  outputs/gene_branch/final_gene_branch_config.json
  outputs/gene_branch/final_gene_branch_metrics.csv
  outputs/gene_branch/final_gene_branch_predictions.csv
  outputs/gene_branch/real_gfv_features_302_subjects.csv
  outputs/gene_branch/final_selected_probe_list.csv
  outputs/gene_branch/final_gene_branch_training_history.csv
  outputs/tables/supplementary/Table_S5_Selected_gene_panel.xlsx
  outputs/tables/supplementary/Table_S6_GE_branch_metrics.xlsx
  manifests/gene_branch_final_split_seed_42.csv
"""

from __future__ import annotations

import json
import pickle
import random
from pathlib import Path

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
from sklearn.impute import SimpleImputer
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    confusion_matrix,
    f1_score,
    precision_score,
    recall_score,
    roc_auc_score,
)
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from torch.utils.data import DataLoader, TensorDataset


# -----------------------------
# Configuration
# -----------------------------

SEED = 42
GFV_DIM = 128
BATCH_SIZE = 16
MAX_EPOCHS = 300
PATIENCE = 30
LR = 1e-4
WEIGHT_DECAY = 5e-3
DROPOUT = 0.60
LABEL_SMOOTHING = 0.05

PROJECT_ROOT = Path(__file__).resolve().parents[3]

DATA_DIR = PROJECT_ROOT / "data"
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "gene_branch"
TABLE_DIR = PROJECT_ROOT / "outputs" / "tables" / "supplementary"
MANIFEST_DIR = PROJECT_ROOT / "manifests"

GE_FILE = DATA_DIR / "gene_expression_data.csv"
META_FILE = DATA_DIR / "Paired_data_metadata.csv"
SELECTED_PANEL_FILE = OUTPUT_DIR / "selected_gene_panel_from_sensitivity.csv"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
TABLE_DIR.mkdir(parents=True, exist_ok=True)
MANIFEST_DIR.mkdir(parents=True, exist_ok=True)


# -----------------------------
# Utilities
# -----------------------------

def set_seed(seed: int = 42) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    torch.cuda.manual_seed_all(seed)


def to_bool(series: pd.Series) -> pd.Series:
    return series.astype(str).str.lower().isin(["true", "1", "yes", "y"])


def count_unique_gene_symbols(gene_series: pd.Series) -> int:
    genes = set()
    for value in gene_series.dropna().astype(str):
        for gene in value.split("||"):
            gene = gene.strip()
            if gene and gene.lower() not in ["nan", "none", ""]:
                genes.add(gene)
    return len(genes)


def safe_auc(y_true: np.ndarray, y_prob: np.ndarray) -> float:
    if len(np.unique(y_true)) < 2:
        return np.nan
    return float(roc_auc_score(y_true, y_prob))


def compute_metrics(y_true: np.ndarray, y_prob: np.ndarray, y_pred: np.ndarray) -> dict:
    tn, fp, fn, tp = confusion_matrix(y_true, y_pred, labels=[0, 1]).ravel()
    specificity = tn / (tn + fp) if (tn + fp) > 0 else np.nan

    return {
        "Subjects, n": int(len(y_true)),
        "Accuracy": accuracy_score(y_true, y_pred),
        "Balanced accuracy": balanced_accuracy_score(y_true, y_pred),
        "AUC": safe_auc(y_true, y_prob),
        "Precision": precision_score(y_true, y_pred, zero_division=0),
        "Recall / AD sensitivity": recall_score(y_true, y_pred, zero_division=0),
        "CN specificity": specificity,
        "F1 score": f1_score(y_true, y_pred, zero_division=0),
        "TN": int(tn),
        "FP": int(fp),
        "FN": int(fn),
        "TP": int(tp),
    }


def style_excel(path: Path) -> None:
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception:
        return

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9D9D9")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_cells in ws.columns:
            values = [str(c.value) if c.value is not None else "" for c in col_cells[:200]]
            width = min(max(len(v) for v in values) + 2, 45)
            ws.column_dimensions[col_cells[0].column_letter].width = width

    wb.save(path)


# -----------------------------
# Model
# -----------------------------

class GeneBranch(nn.Module):
    def __init__(self, input_dim: int, gfv_dim: int = 128, dropout: float = 0.60):
        super().__init__()

        self.encoder = nn.Sequential(
            nn.LayerNorm(input_dim),
            nn.Dropout(dropout),
            nn.Linear(input_dim, gfv_dim),
            nn.LayerNorm(gfv_dim),
            nn.ReLU(),
            nn.Dropout(dropout),
        )

        self.classifier = nn.Linear(gfv_dim, 2)

    def forward(self, x):
        gfv = self.encoder(x)
        logits = self.classifier(gfv)
        return logits, gfv


def predict(model: nn.Module, x_np: np.ndarray, device: torch.device) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    model.eval()
    probs_all = []
    preds_all = []
    gfv_all = []

    loader = DataLoader(
        TensorDataset(torch.tensor(x_np, dtype=torch.float32)),
        batch_size=128,
        shuffle=False,
    )

    with torch.no_grad():
        for (xb,) in loader:
            xb = xb.to(device)
            logits, gfv = model(xb)
            probs = torch.softmax(logits, dim=1)[:, 1]
            preds = torch.argmax(logits, dim=1)

            probs_all.append(probs.cpu().numpy())
            preds_all.append(preds.cpu().numpy())
            gfv_all.append(gfv.cpu().numpy())

    return np.concatenate(probs_all), np.concatenate(preds_all), np.vstack(gfv_all)


def evaluate_loss_auc(model, x_np, y_np, criterion, device):
    model.eval()
    losses = []
    probs = []
    preds = []

    loader = DataLoader(
        TensorDataset(
            torch.tensor(x_np, dtype=torch.float32),
            torch.tensor(y_np, dtype=torch.long),
        ),
        batch_size=128,
        shuffle=False,
    )

    with torch.no_grad():
        for xb, yb in loader:
            xb = xb.to(device)
            yb = yb.to(device)

            logits, _ = model(xb)
            loss = criterion(logits, yb)
            prob = torch.softmax(logits, dim=1)[:, 1]
            pred = torch.argmax(logits, dim=1)

            losses.append(loss.item() * len(yb))
            probs.append(prob.cpu().numpy())
            preds.append(pred.cpu().numpy())

    probs = np.concatenate(probs)
    preds = np.concatenate(preds)

    return {
        "loss": float(np.sum(losses) / len(y_np)),
        "auc": safe_auc(y_np, probs),
        "balanced_accuracy": balanced_accuracy_score(y_np, preds),
        "accuracy": accuracy_score(y_np, preds),
        "f1": f1_score(y_np, preds, zero_division=0),
    }


# -----------------------------
# Data loading
# -----------------------------

def make_unique_columns(cols):
    """
    Recreate pandas-style duplicate column names:
    ADNI2, ADNI2, ADNI2 -> ADNI2, ADNI2.1, ADNI2.2
    """
    clean_cols = []
    counts = {}

    for c in cols:
        c = "" if pd.isna(c) else str(c).strip()

        if c in counts:
            counts[c] += 1
            clean_cols.append(f"{c}.{counts[c]}")
        else:
            counts[c] = 0
            clean_cols.append(c)

    return clean_cols


def load_expression_matrix() -> tuple[pd.DataFrame, str, str, str]:
    """
    Load gene_expression_data.csv using the same logic as the already-working
    DEG and gene-panel sensitivity scripts.

    Important:
    - The CSV header already contains GE sample columns such as ADNI2, ADNI2.1, ADNIGO.1.
    - The row where first column == ProbeSet marks the start of the expression matrix.
    - We should not use the ProbeSet row as the CSV header.
    """
    print("Reading gene-expression file...")
    ge = pd.read_csv(GE_FILE, low_memory=False)

    first_col = ge.iloc[:, 0].astype(str).str.strip()
    matches = ge.index[first_col == "ProbeSet"].tolist()

    if not matches:
        raise ValueError("Could not find row labeled 'ProbeSet' in the first column.")

    probeset_row = matches[0]
    expr = ge.iloc[probeset_row + 1:].copy()

    probe_col = expr.columns[0]
    locus_col = expr.columns[1]
    symbol_col = expr.columns[2]

    expr[probe_col] = expr[probe_col].astype(str).str.strip()

    print(f"Expression features loaded: {expr.shape[0]:,}")
    print(f"Expression sample columns available: {expr.shape[1] - 3:,}")
    print(f"Example GE columns from CSV header: {list(ge.columns[3:13])}")

    return expr, probe_col, locus_col, symbol_col


def prepare_dataset():
    selected_panel = pd.read_csv(SELECTED_PANEL_FILE)
    if "Probe ID / feature ID" not in selected_panel.columns:
        raise ValueError("Selected panel file must contain 'Probe ID / feature ID'.")

    selected_panel["Probe ID / feature ID"] = selected_panel["Probe ID / feature ID"].astype(str)
    panel_probes = selected_panel["Probe ID / feature ID"].tolist()

    expr, probe_col, locus_col, symbol_col = load_expression_matrix()

    meta = pd.read_csv(META_FILE)

    required_cols = ["subject_id", "ge_sample_column", "diagnosis", "diagnosis_binary", "ge_column_found"]
    missing_cols = [c for c in required_cols if c not in meta.columns]
    if missing_cols:
        raise ValueError(f"Missing columns in metadata: {missing_cols}")

    meta = meta.copy()
    meta = meta[to_bool(meta["ge_column_found"])]
    meta = meta[meta["diagnosis"].astype(str).isin(["AD", "CN"])]
    meta["diagnosis_binary"] = meta["diagnosis_binary"].astype(int)

    sample_cols = meta["ge_sample_column"].astype(str).tolist()
    missing_sample_cols = [c for c in sample_cols if c not in expr.columns]
    if missing_sample_cols:
        raise ValueError(
            f"{len(missing_sample_cols)} metadata GE sample columns are missing from expression matrix. "
            f"Examples: {missing_sample_cols[:10]}"
        )

    expr_panel = expr[expr[probe_col].astype(str).isin(panel_probes)].copy()
    expr_panel = expr_panel.drop_duplicates(subset=[probe_col], keep="first")
    expr_panel = expr_panel.set_index(probe_col)

    available_probes = [p for p in panel_probes if p in expr_panel.index]
    missing_probes = [p for p in panel_probes if p not in expr_panel.index]

    if len(available_probes) == 0:
        raise ValueError("None of the selected probes were found in the expression matrix.")

    expr_panel = expr_panel.loc[available_probes]

    x_df = expr_panel[sample_cols].T
    x_df.index = meta.index
    x_df = x_df.apply(pd.to_numeric, errors="coerce")

    y = meta["diagnosis_binary"].to_numpy(dtype=int)

    selected_panel_used = selected_panel[
        selected_panel["Probe ID / feature ID"].astype(str).isin(available_probes)
    ].copy()

    selected_panel_used["Selected probe panel label"] = "Top 2000 limma-ranked probes"
    selected_panel_used["Selected probe panel size"] = len(available_probes)

    if "Gene symbol" in selected_panel_used.columns:
        unique_gene_count = count_unique_gene_symbols(selected_panel_used["Gene symbol"])
    else:
        unique_gene_count = np.nan

    selected_panel_used["Approximate unique gene symbols in selected panel"] = unique_gene_count

    print("\n========== FINAL GE BRANCH INPUT ==========")
    print(f"Measured GE subjects: {len(meta):,}")
    print(f"CN: {(y == 0).sum():,}")
    print(f"AD: {(y == 1).sum():,}")
    print(f"Selected probes requested: {len(panel_probes):,}")
    print(f"Actual probes used: {len(available_probes):,}")
    print(f"Missing selected probes: {len(missing_probes):,}")
    print(f"Approximate unique gene symbols: {unique_gene_count:,}")
    print("==========================================\n")

    return x_df, y, meta.reset_index(drop=True), selected_panel_used, unique_gene_count


# -----------------------------
# Main training
# -----------------------------

def main():
    set_seed(SEED)

    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")

    print("========== TRAIN FINAL GE BRANCH ==========")
    print(f"Project root: {PROJECT_ROOT}")
    print(f"Device: {device}")
    if torch.cuda.is_available():
        print(f"GPU: {torch.cuda.get_device_name(0)}")
    print("==========================================\n")

    x_df, y, meta, selected_panel_used, unique_gene_count = prepare_dataset()
    n_subjects, input_dim = x_df.shape

    # ---------------------------------------------------------
    # Canonical split rule
    # ---------------------------------------------------------
    # The MRI branch defines the canonical train/validation/test split.
    # For paired GE+MRI subjects, inherit the MRI split.
    # For GE-only subjects without paired MRI, assign to Train.
    # This keeps GE, MRI, and XIG-CVAE splits consistent and reproducible.
    MRI_SPLIT_PATH = MANIFEST_DIR / "mri_branch_final_split_seed_42.csv"

    if not MRI_SPLIT_PATH.exists():
        raise FileNotFoundError(
            f"Canonical MRI split manifest not found: {MRI_SPLIT_PATH}"
        )

    mri_manifest = pd.read_csv(MRI_SPLIT_PATH, dtype={"subject_id": str})
    if "subject_id" not in mri_manifest.columns or "split" not in mri_manifest.columns:
        raise ValueError("MRI split manifest must contain subject_id and split columns.")

    mri_split = (
        mri_manifest[["subject_id", "split"]]
        .drop_duplicates(subset=["subject_id"])
        .rename(columns={"split": "canonical_mri_split"})
        .copy()
    )

    manifest = meta.copy()
    manifest["subject_id"] = manifest["subject_id"].astype(str)
    mri_split["subject_id"] = mri_split["subject_id"].astype(str)

    manifest = manifest.merge(mri_split, on="subject_id", how="left", sort=False)

    manifest["has_paired_mri"] = manifest["canonical_mri_split"].notna()

    # GE-only subjects are used only for training the GE encoder.
    manifest["split"] = manifest["canonical_mri_split"].fillna("Train")

    allowed_splits = {"Train", "Validation", "Test"}
    observed_splits = set(manifest["split"].unique())
    if not observed_splits.issubset(allowed_splits):
        raise ValueError(f"Unexpected split labels found: {observed_splits}")

    split = manifest["split"].to_numpy()

    train_idx = np.where(split == "Train")[0]
    val_idx = np.where(split == "Validation")[0]
    test_idx = np.where(split == "Test")[0]

    if len(train_idx) == 0 or len(val_idx) == 0 or len(test_idx) == 0:
        raise ValueError(
            f"Empty split detected: Train={len(train_idx)}, "
            f"Validation={len(val_idx)}, Test={len(test_idx)}"
        )

    manifest_path = MANIFEST_DIR / "gene_branch_final_split_seed_42.csv"
    manifest.to_csv(manifest_path, index=False)

    print("Canonical split source:")
    print(MRI_SPLIT_PATH)
    print()

    print("GE subjects with paired MRI:", int(manifest["has_paired_mri"].sum()))
    print("GE-only subjects assigned to Train:", int((~manifest["has_paired_mri"]).sum()))
    print()

    print("Split counts:")
    print(manifest.groupby(["split", "diagnosis"]).size().unstack(fill_value=0))
    print()

    print("Paired-only split counts:")
    print(
        manifest.loc[manifest["has_paired_mri"]]
        .groupby(["split", "diagnosis"])
        .size()
        .unstack(fill_value=0)
    )
    print()

    preprocess = Pipeline(
        steps=[
            ("imputer", SimpleImputer(strategy="mean")),
            ("scaler", StandardScaler()),
        ]
    )

    x_train = preprocess.fit_transform(x_df.iloc[train_idx])
    x_val = preprocess.transform(x_df.iloc[val_idx])
    x_test = preprocess.transform(x_df.iloc[test_idx])
    x_all = preprocess.transform(x_df)

    y_train = y[train_idx]
    y_val = y[val_idx]
    y_test = y[test_idx]

    class_counts = np.bincount(y_train, minlength=2)
    class_weights = len(y_train) / (2.0 * np.maximum(class_counts, 1))
    class_weights_t = torch.tensor(class_weights, dtype=torch.float32).to(device)

    model = GeneBranch(input_dim=input_dim, gfv_dim=GFV_DIM, dropout=DROPOUT).to(device)

    criterion = nn.CrossEntropyLoss(
        weight=class_weights_t,
        label_smoothing=LABEL_SMOOTHING,
    )

    optimizer = torch.optim.Adam(
        model.parameters(),
        lr=LR,
        weight_decay=WEIGHT_DECAY,
    )

    train_loader = DataLoader(
        TensorDataset(
            torch.tensor(x_train, dtype=torch.float32),
            torch.tensor(y_train, dtype=torch.long),
        ),
        batch_size=BATCH_SIZE,
        shuffle=True,
    )

    best_state = None
    best_epoch = 0
    best_score = (-np.inf, -np.inf, np.inf)
    patience_counter = 0
    history = []

    for epoch in range(1, MAX_EPOCHS + 1):
        model.train()
        train_losses = []

        for xb, yb in train_loader:
            xb = xb.to(device)
            yb = yb.to(device)

            optimizer.zero_grad()
            logits, _ = model(xb)
            loss = criterion(logits, yb)
            loss.backward()
            optimizer.step()

            train_losses.append(loss.item() * len(yb))

        train_eval = evaluate_loss_auc(model, x_train, y_train, criterion, device)
        val_eval = evaluate_loss_auc(model, x_val, y_val, criterion, device)

        row = {
            "epoch": epoch,
            "train_loss": train_eval["loss"],
            "train_auc": train_eval["auc"],
            "train_balanced_accuracy": train_eval["balanced_accuracy"],
            "train_accuracy": train_eval["accuracy"],
            "train_f1": train_eval["f1"],
            "validation_loss": val_eval["loss"],
            "validation_auc": val_eval["auc"],
            "validation_balanced_accuracy": val_eval["balanced_accuracy"],
            "validation_accuracy": val_eval["accuracy"],
            "validation_f1": val_eval["f1"],
        }
        history.append(row)

        current_score = (
            val_eval["auc"] if not np.isnan(val_eval["auc"]) else -np.inf,
            val_eval["balanced_accuracy"],
            -val_eval["loss"],
        )

        if current_score > best_score:
            best_score = current_score
            best_epoch = epoch
            best_state = {
                "model_state_dict": model.state_dict(),
                "input_dim": input_dim,
                "gfv_dim": GFV_DIM,
                "dropout": DROPOUT,
                "best_epoch": best_epoch,
                "seed": SEED,
            }
            patience_counter = 0
        else:
            patience_counter += 1

        if epoch == 1 or epoch % 10 == 0:
            print(
                f"Epoch {epoch:03d} | "
                f"train loss {train_eval['loss']:.4f} | "
                f"val loss {val_eval['loss']:.4f} | "
                f"val AUC {val_eval['auc']:.3f} | "
                f"val balanced acc {val_eval['balanced_accuracy']:.3f}"
            )

        if patience_counter >= PATIENCE:
            print(f"Early stopping at epoch {epoch}. Best epoch: {best_epoch}")
            break

    if best_state is None:
        raise RuntimeError("Training failed: no best model state was recorded.")

    model.load_state_dict(best_state["model_state_dict"])

    train_prob, train_pred, _ = predict(model, x_train, device)
    val_prob, val_pred, _ = predict(model, x_val, device)
    test_prob, test_pred, _ = predict(model, x_test, device)
    all_prob, all_pred, all_gfv = predict(model, x_all, device)

    metrics_rows = []

    for split_name, indices, yy, prob, pred in [
        ("Train", train_idx, y_train, train_prob, train_pred),
        ("Validation", val_idx, y_val, val_prob, val_pred),
        ("Test", test_idx, y_test, test_prob, test_pred),
    ]:
        m = compute_metrics(yy, prob, pred)
        m.update(
            {
                "Model": "Final GE branch",
                "Input representation": "Top 2000 limma-ranked probes -> GFV128",
                "Actual probes used": input_dim,
                "Approximate unique gene symbols": unique_gene_count,
                "GFV dimension": GFV_DIM,
                "Split": split_name,
                "Best epoch": best_epoch,
            }
        )
        metrics_rows.append(m)

    metrics_df = pd.DataFrame(metrics_rows)

    ordered_cols = [
        "Model",
        "Input representation",
        "Actual probes used",
        "Approximate unique gene symbols",
        "GFV dimension",
        "Split",
        "Subjects, n",
        "Best epoch",
        "Accuracy",
        "Balanced accuracy",
        "AUC",
        "Precision",
        "Recall / AD sensitivity",
        "CN specificity",
        "F1 score",
        "TN",
        "FP",
        "FN",
        "TP",
    ]

    metrics_df = metrics_df[ordered_cols]

    predictions_df = meta.copy()
    predictions_df["split"] = split
    predictions_df["true_label"] = y
    predictions_df["predicted_AD_probability"] = all_prob
    predictions_df["predicted_label"] = all_pred
    predictions_df["predicted_diagnosis"] = np.where(all_pred == 1, "AD", "CN")

    gfv_cols = [f"GFV{j+1:03d}" for j in range(GFV_DIM)]
    gfv_df = pd.concat(
        [
            meta[["subject_id", "bids_id", "rid", "ge_sample_column", "diagnosis", "diagnosis_binary"]].copy(),
            pd.DataFrame({"split": split}),
            pd.DataFrame(all_gfv, columns=gfv_cols),
        ],
        axis=1,
    )

    history_df = pd.DataFrame(history)

    model_path = OUTPUT_DIR / "final_gene_branch_model.pt"
    scaler_path = OUTPUT_DIR / "final_gene_branch_scaler.pkl"
    config_path = OUTPUT_DIR / "final_gene_branch_config.json"
    metrics_path = OUTPUT_DIR / "final_gene_branch_metrics.csv"
    predictions_path = OUTPUT_DIR / "final_gene_branch_predictions.csv"
    gfv_path = OUTPUT_DIR / "real_gfv_features_302_subjects.csv"
    selected_probe_path = OUTPUT_DIR / "final_selected_probe_list.csv"
    history_path = OUTPUT_DIR / "final_gene_branch_training_history.csv"

    torch.save(best_state, model_path)

    with open(scaler_path, "wb") as f:
        pickle.dump(preprocess, f)

    config = {
        "model": "Final GE branch",
        "seed": SEED,
        "input_panel": "Top 2000 limma-ranked probes",
        "actual_probes_used": int(input_dim),
        "approximate_unique_gene_symbols": int(unique_gene_count) if not pd.isna(unique_gene_count) else None,
        "gfv_dimension": GFV_DIM,
        "dropout": DROPOUT,
        "batch_size": BATCH_SIZE,
        "max_epochs": MAX_EPOCHS,
        "early_stopping_patience": PATIENCE,
        "learning_rate": LR,
        "weight_decay": WEIGHT_DECAY,
        "label_smoothing": LABEL_SMOOTHING,
        "split": "MRI-branch canonical split for paired subjects; GE-only subjects assigned to Train",
        "split_source": str(MANIFEST_DIR / "mri_branch_final_split_seed_42.csv"),
        "best_epoch": int(best_epoch),
        "subjects_total": int(n_subjects),
        "train_n": int(len(train_idx)),
        "validation_n": int(len(val_idx)),
        "test_n": int(len(test_idx)),
    }

    with open(config_path, "w") as f:
        json.dump(config, f, indent=2)

    metrics_df.to_csv(metrics_path, index=False)
    predictions_df.to_csv(predictions_path, index=False)
    gfv_df.to_csv(gfv_path, index=False)
    selected_panel_used.to_csv(selected_probe_path, index=False)
    history_df.to_csv(history_path, index=False)

    table_s5_path = TABLE_DIR / "Table_S5_Selected_gene_panel.xlsx"
    table_s6_path = TABLE_DIR / "Table_S6_GE_branch_metrics.xlsx"

    with pd.ExcelWriter(table_s5_path, engine="openpyxl") as writer:
        selected_panel_used.to_excel(writer, index=False, sheet_name="Selected gene panel")

    with pd.ExcelWriter(table_s6_path, engine="openpyxl") as writer:
        metrics_df.to_excel(writer, index=False, sheet_name="GE branch metrics")

    style_excel(table_s5_path)
    style_excel(table_s6_path)

    print("\n========== FINAL GE BRANCH METRICS ==========")
    print(f"Actual probes used: {input_dim}")
    print(f"Approximate unique gene symbols: {unique_gene_count}")
    print(f"GFV dimension: {GFV_DIM}")
    print(f"Best epoch: {best_epoch}")
    print()
    print(metrics_df.to_string(index=False))
    print("============================================\n")

    print("Saved:")
    for p in [
        model_path,
        scaler_path,
        config_path,
        metrics_path,
        predictions_path,
        gfv_path,
        selected_probe_path,
        history_path,
        table_s5_path,
        table_s6_path,
        manifest_path,
    ]:
        print(p)


if __name__ == "__main__":
    main()
