"""
Reusable Excel styling utilities for manuscript tables.

Main tables should use stronger journal-style formatting.
Supplementary tables should remain clean, readable, and filterable.
"""

from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def style_supplementary_table(
    path,
    sheet_name,
    freeze_panes="A2",
    header_fill="D9D9D9",
    max_width=38,
):
    path = Path(path)
    wb = load_workbook(path)
    ws = wb[sheet_name]

    header = PatternFill("solid", fgColor=header_fill)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Times New Roman", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    for cell in ws[1]:
        cell.font = Font(name="Times New Roman", size=11, bold=True)
        cell.fill = header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)

        values = []
        for row_idx in range(1, min(ws.max_row, 200) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                values.append(str(value))

        if values:
            width = min(max(max(len(v) for v in values) + 2, 12), max_width)
        else:
            width = 12

        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = freeze_panes
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
